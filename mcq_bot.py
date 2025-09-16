import os, logging, tempfile, json
import fitz  # PyMuPDF for PDF
import docx  # python-docx for Word
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    ApplicationBuilder, CommandHandler, MessageHandler,
    CallbackQueryHandler, filters, ContextTypes
)

# ========= CONFIG =========
BOT_TOKEN = os.getenv("BOT_TOKEN")
ADMIN_ID = int(os.getenv("ADMIN_ID", "0"))

logging.basicConfig(level=logging.INFO)
user_sessions = {}
LANGUAGE = "bn"  # default Bengali
NUM_QUESTIONS = 5
DB_FILE = "mcq_bank.json"

# -------- LOAD / SAVE BANK --------
def load_bank():
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            return {}
    return {}

def save_bank(data):
    with open(DB_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

mcq_bank = load_bank()

# -------- BASIC MCQ GENERATOR --------
def extract_text_from_pdf(path):
    text = ""
    with fitz.open(path) as doc:
        for page in doc:
            text += page.get_text()
    return text

def extract_text_from_docx(path):
    doc = docx.Document(path)
    return "\n".join([p.text for p in doc.paragraphs])

def generate_basic_mcqs(text, num=5, lang="bn"):
    lines = [l.strip() for l in text.split("\n") if len(l.strip()) > 10]
    mcqs = []
    for i, line in enumerate(lines[:max(1, num)]):
        q_lang = "প্রশ্ন" if lang == "bn" else "Question"
        expl_lang = "সঠিক উত্তর এই তথ্য থেকে পাওয়া যায়।" if lang == "bn" else "The correct answer comes from this line."
        words = line.split()
        first_word = words[0] if len(words) > 0 else "Term"
        # create simple distractors by rotating words or using placeholders
        opt1 = first_word
        opt2 = words[1] if len(words) > 1 else "Option A"
        opt3 = words[2] if len(words) > 2 else "Option B"
        opt4 = "Option C"
        mcqs.append({
            "question": f"{q_lang}: {first_word} সম্পর্কে কোনটি সঠিক?" if lang=="bn"
                        else f"{q_lang}: Which of the following about '{first_word}' is correct?",
            "options": [opt1, opt2, opt3, opt4],
            "answer": 0,
            "explanation": f"{expl_lang} → {line}"
        })
    return mcqs

# -------- TELEGRAM HANDLERS --------
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("👋 Send me a PDF/Word/TXT and I'll generate MCQs. Use /setlang en or /setlang bn and /setnum N to change settings. Admin: use /admin")

async def handle_file(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if doc is None:
        await update.message.reply_text("Please send a document (PDF/DOCX/TXT).")
        return
    file = await doc.get_file()
    suffix = ""
    if doc.file_name:
        suffix = os.path.splitext(doc.file_name)[1].lower()
    # save temp file
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        await file.download_to_drive(tmp.name)
        # extract
        if tmp.name.endswith(".pdf"):
            text = extract_text_from_pdf(tmp.name)
        elif tmp.name.endswith(".docx"):
            text = extract_text_from_docx(tmp.name)
        else:
            try:
                with open(tmp.name, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            except:
                text = ""
    if not text:
        await update.message.reply_text("Could not extract text from the document. Try a plain TXT or a different PDF.")
        return
    mcqs = generate_basic_mcqs(text, NUM_QUESTIONS, LANGUAGE)
    user_sessions[update.effective_user.id] = {"mcqs": mcqs}
    await update.message.reply_text(f"✅ Generated {len(mcqs)} MCQs. Start answering below.")
    await send_question(update, context, 0)

async def send_question(update, context, index):
    chat_id = update.effective_chat.id
    session = user_sessions.get(update.effective_user.id, {})
    mcqs = session.get("mcqs", [])
    if index >= len(mcqs):
        await context.bot.send_message(chat_id, "🎉 Done! /start to try again.")
        return
    q = mcqs[index]
    keyboard = [[InlineKeyboardButton(opt, callback_data=f"{index}:{i}") for i,opt in enumerate(q["options"])]]
    await context.bot.send_message(chat_id, f"❓ {q['question']}", reply_markup=InlineKeyboardMarkup(keyboard))

async def handle_answer(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    if not query:
        return
    await query.answer()
    try:
        index, chosen = map(int, query.data.split(":"))
    except:
        await query.edit_message_text("⚠️ Invalid response.")
        return
    sess = user_sessions.get(update.effective_user.id, {})
    mcqs = sess.get("mcqs", [])
    if index < 0 or index >= len(mcqs):
        await query.edit_message_text("⚠️ Question index error.")
        return
    q = mcqs[index]
    correct = q.get("answer", 0)
    result = "✅ Correct!" if chosen == correct else f"❌ Wrong! Correct: {q['options'][correct]}"
    explanation = f"📖 {q.get('explanation','')}"
    await query.edit_message_text(f"❓ {q['question']}\n\n{result}\n{explanation}")
    # send next
    await send_question(update, context, index+1)

# -------- ADMIN / UTIL COMMANDS --------
async def admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return await update.message.reply_text("❌ Permission denied.")
    text = (
        "🛠 Admin commands:\n"
        "/backup - save current sessions to mcq_bank.json and send file\n"
        "/dumpbank - send the stored mcq_bank.json file\n"
        "/setlang en|bn - set language\n"
        "/setnum N - set number of MCQs\n    "
    )
    await update.message.reply_text(text)

async def backup_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return await update.message.reply_text("❌ Permission denied.")
    # merge user_sessions into mcq_bank using user id as key
    changed = False
    for uid, sess in user_sessions.items():
        qlist = sess.get("mcqs", [])
        if qlist:
            mcq_bank[str(uid)] = qlist
            changed = True
    if changed:
        save_bank(mcq_bank)
    await update.message.reply_text("✅ Backup saved to mcq_bank.json. Sending file...")
    if os.path.exists(DB_FILE):
        await context.bot.send_document(update.effective_chat.id, document=open(DB_FILE, "rb"))
    else:
        await update.message.reply_text("No bank file present.")

async def dumpbank_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID:
        return await update.message.reply_text("❌ Permission denied.")
    if os.path.exists(DB_FILE):
        await context.bot.send_document(update.effective_chat.id, document=open(DB_FILE, "rb"))
    else:
        await update.message.reply_text("No mcq_bank.json file found.")

# -------- LANGUAGE & SETTINGS --------
async def set_lang(update: Update, context: ContextTypes.DEFAULT_TYPE):
    global LANGUAGE
    if context.args and context.args[0] in ["en", "bn"]:
        LANGUAGE = context.args[0]
        msg = "✅ Language set to English" if LANGUAGE == "en" else "✅ ভাষা বাংলায় সেট হয়েছে"
        await update.message.reply_text(msg)
    else:
        await update.message.reply_text("Usage: /setlang en  OR  /setlang bn")

async def set_num(update: Update, context: ContextTypes.DEFAULT_TYPE):
    global NUM_QUESTIONS
    try:
        NUM_QUESTIONS = int(context.args[0])
        await update.message.reply_text(f"✅ Number of questions set to {NUM_QUESTIONS}")
    except:
        await update.message.reply_text("⚠️ Usage: /setnum 10")

# -------- BOT RUN --------
def main():
    if not BOT_TOKEN:
        logging.error("BOT_TOKEN not set in environment variables.")
        return
    app = ApplicationBuilder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("setlang", set_lang))
    app.add_handler(CommandHandler("setnum", set_num))
    app.add_handler(CommandHandler("admin", admin_command))
    app.add_handler(CommandHandler("backup", backup_command))
    app.add_handler(CommandHandler("dumpbank", dumpbank_command))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_file))
    app.add_handler(CallbackQueryHandler(handle_answer))
    logging.info("Bot starting...")
    app.run_polling()

if __name__ == "__main__":
    main()
